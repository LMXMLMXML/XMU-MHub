# -*- coding: utf-8 -*-
"""DEMO.md → 厦大统一门户-功能演示文档.docx

为什么要用脚本而不是手写：DEMO.md 是唯一的内容源，改了它重跑一次，
Word 版就跟着更新，不会出现两份内容对不上。

只用 python-docx（构建期工具，不进软件本体——软件本体仍然是零依赖）。

用法：
    python tools/make_demo_docx.py
产物：
    厦大统一门户-功能演示文档.docx
"""
from __future__ import annotations

import datetime as _dt
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "DEMO.md"
DST = ROOT / "厦大统一门户-功能演示文档.docx"
LOGO = ROOT / "web" / "assets" / "logo" / "logo-192.png"

# 取自 web/styles.css 的品牌色，保证文档和软件看起来是一家的
BRAND = "0D2F6E"
BRAND_MID = "134C92"
ACCENT = "0C81BD"
GREEN = "4F9E3F"
SOFT = "E3F2FB"
SOFT2 = "F4F7FB"
GREEN_SOFT = "EAF6E6"
BORDER = "D5DEEB"
DIM = "55678A"
MUTE = "8A99B0"
TEXT = "0F1B33"

FONT = "微软雅黑"
MONO = "Consolas"
WIDTH_CM = 16.6          # A4 去掉左右 2.2cm 边距后的正文宽度

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


# ---------------------------------------------------------------- 小工具
def dw(text: str) -> int:
    """显示宽度：汉字算 2（算表格列宽用）"""
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def font_of(obj, name: str, size: float | None = None, bold: bool | None = None,
            color: str | None = None, italic: bool | None = None) -> None:
    """同时设置西文与中文（eastAsia）字体，否则 Word 会用宋体兜底"""
    f = obj.font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = RGBColor.from_string(color)
    rpr = obj.element.get_or_add_rPr() if hasattr(obj.element, "get_or_add_rPr") else obj.element.rPr
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)
    rf.set(qn("w:eastAsia"), FONT if name != MONO else FONT)


def shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


# OOXML 里子元素的顺序是写死的，顺序不对 Word 会弹"需要修复"。
# 下面按 schema 顺序插入，而不是无脑 append。
ORDER = {
    "w:pPr": ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr",
              "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
              "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
              "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
              "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
              "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"],
    "w:tcPr": ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders",
               "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
               "w:vAlign", "w:hideMark"],
    "w:tblPr": ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
                "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
                "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout",
                "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"],
}


def put(parent, tag: str) -> object:
    """按 schema 顺序插入（已存在则直接返回）"""
    got = parent.find(qn(tag))
    if got is not None:
        return got
    el = OxmlElement(tag)
    order = ORDER["w:" + parent.tag.split("}")[-1]]
    idx = order.index(tag)
    for child in parent:
        cname = "w:" + child.tag.split("}")[-1]
        if cname in order and order.index(cname) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def para_border(p, edges: str, color: str = BORDER, size: int = 8, space: int = 6) -> None:
    pBdr = put(p._p.get_or_add_pPr(), "w:pBdr")
    for edge in edges.split():
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), color)
        pBdr.append(e)


def para_shade(p, fill: str) -> None:
    shade(put(p._p.get_or_add_pPr(), "w:shd"), fill)


def cell_shade(cell, fill: str) -> None:
    shade(put(cell._tc.get_or_add_tcPr(), "w:shd"), fill)


def cell_margins(table, top=60, bottom=60, left=110, right=110) -> None:
    mar = put(table._tbl.tblPr, "w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + name)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)


def table_borders(table, color: str = BORDER, inner: str | None = None) -> None:
    borders = put(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6" if edge.startswith("inside") else "8")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), (inner or color) if edge.startswith("inside") else color)
        borders.append(e)


def fixed_layout(table, widths_cm: list[float]) -> None:
    layout = put(table._tbl.tblPr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)
    for i, w in enumerate(widths_cm):
        if i < len(table.columns):
            table.columns[i].width = Cm(w)


def repeat_header(table) -> None:
    trPr = table.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def col_widths(rows: list[list[str]], total: float = WIDTH_CM) -> list[float]:
    cols = max(len(r) for r in rows)
    weights = []
    for c in range(cols):
        mx = max([dw(r[c]) for r in rows if c < len(r)] or [4])
        weights.append(max(4, min(mx, 46)))
    s = float(sum(weights))
    out = [total * w / s for w in weights]
    out = [max(1.9, w) for w in out]
    k = total / sum(out)
    return [w * k for w in out]


# ------------------------------------------------- 行内 **粗体** / `代码` / *斜体*
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+\*)")


def add_inline(p, text: str, size: float, color: str = TEXT, bold: bool = False,
               code_color: str = BRAND_MID, italic: bool = False) -> None:
    """把 **粗体**、`代码`、*斜体* 变成真正的 Word 格式；可以嵌套（**`代码`**）"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(p, part[2:-2], size, color, True, code_color, italic)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1])
            font_of(r, MONO, size - 0.5, bold, code_color, italic=italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(p, part[1:-1], size, color, bold, code_color, True)
        else:
            r = p.add_run(part)
            font_of(r, FONT, size, bold, color, italic=italic)


LABEL_RE = re.compile(r"^\*\*([^*]{1,16})\*\*：(.*)$", re.S)
LABEL_RE2 = re.compile(r"^([^：*]{1,15})：(.*)$", re.S)
LABELS = {"操作", "屏幕表现", "说明", "证据", "指着 ⚠ 角标说", "操作顺序", "备用",
          "演示两个细节", "再补两刀", "再点一条打不开的"}


def split_label(text: str) -> tuple[str | None, str]:
    """拆出「操作：」这种小标题（可能被 ** 包着）"""
    for rx in (LABEL_RE, LABEL_RE2):
        m = rx.match(text)
        if m:
            label = m.group(1).strip()
            if label in LABELS or (len(label) <= 14 and not set(label) & set("。？！，、；")):
                return label, m.group(2)
            return None, text
    return None, text


# ---------------------------------------------------------------- 读 DEMO.md
def join_lines(buf: list[str]) -> str:
    """把 markdown 里被硬折行的同一段合并：中文直接接，英文之间补空格"""
    out = ""
    for ln in buf:
        ln = ln.strip()
        if not ln:
            continue
        if out and out[-1].isascii() and out[-1].isalnum() and ln[0].isascii() and ln[0].isalnum():
            out += " "
        out += ln
    return out


def parse(md: str) -> list[tuple]:
    blocks: list[tuple] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", buf))
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue

        if line.strip() == "---":
            blocks.append(("hr", None))
            i += 1
            continue

        m = re.match(r"^(#+)\s+(.*)$", line)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2).strip()))
            i += 1
            continue

        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip(">").strip())
                i += 1
            blocks.append(("quote", [b for b in buf if b]))
            continue

        if re.match(r"^\s*[-*]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("ul", items))
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(lines[i].strip())
                i += 1
            blocks.append(("ol", items))
            continue

        if not line.strip():
            i += 1
            continue

        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^\s*(#|>|\||-{3,}|\*\s|\d+\.\s|```)", lines[i]):
            buf.append(lines[i])
            i += 1
        blocks.append(("p", join_lines(buf)))
    return blocks


# ---------------------------------------------------------------- 造 Word
def build_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    font_of(normal, FONT, 10.5, False, TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28

    def new(name: str, size: float, color: str, bold: bool = False,
            base: str = "Normal", ea_font: str | None = None) -> object:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
        st.quick_style = True
        font_of(st, ea_font or FONT, size, bold, color)
        return st

    new("DemoTitle", 24, BRAND, True).paragraph_format.space_after = Pt(2)
    new("DemoSub", 10.5, DIM).paragraph_format.space_after = Pt(10)
    h1 = new("DemoH1", 15, BRAND, True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h2 = new("DemoH2", 12, BRAND_MID, True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    new("DemoH3", 11, ACCENT, True)

    body = new("DemoBody", 10.5, TEXT)
    body.paragraph_format.space_after = Pt(7)

    quote = new("DemoQuote", 10.5, BRAND_MID)
    quote.paragraph_format.left_indent = Cm(0.5)
    quote.paragraph_format.right_indent = Cm(0.3)
    quote.paragraph_format.space_before = Pt(2)
    quote.paragraph_format.space_after = Pt(2)

    code = new("DemoCode", 9, "1F3A5F", ea_font=MONO)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.line_spacing = 1.15

    ev = new("DemoEvidence", 9, DIM)
    ev.paragraph_format.left_indent = Cm(0.35)
    ev.paragraph_format.space_after = Pt(7)

    ul = new("DemoBullet", 10.5, TEXT)
    ul.paragraph_format.left_indent = Cm(0.75)
    ul.paragraph_format.first_line_indent = Cm(-0.35)
    ul.paragraph_format.space_after = Pt(5)

    ol = new("DemoNumber", 10.5, TEXT)
    ol.paragraph_format.left_indent = Cm(0.75)
    ol.paragraph_format.first_line_indent = Cm(-0.35)
    ol.paragraph_format.space_after = Pt(5)


def build_header_footer(doc: Document, section) -> None:
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)

    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(2)
    if LOGO.exists():
        r = hp.add_run()
        r.add_picture(str(LOGO), height=Cm(0.72))
    hp.add_run("\t")
    r2 = hp.add_run("厦大统一门户 · 功能演示文档")
    font_of(r2, FONT, 8.5, False, MUTE)
    para_border(hp, "bottom", BORDER, 6, 4)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    font_of(r, FONT, 8.5, False, MUTE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fr = OxmlElement("w:r")
    ft = OxmlElement("w:t")
    ft.text = "1"
    fr.append(ft)
    fld.append(fr)
    fp._p.append(fld)
    r2 = fp.add_run(" 页　·　数据版本 320 条 / 39 组　·　生成于 %s"
                    % _dt.date.today().isoformat())
    font_of(r2, FONT, 8.5, False, MUTE)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=cols)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t)

    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.18
            if ri == 0:
                add_inline(p, text, 9.5, "FFFFFF", bold=True, code_color="DCE9FA")
                cell_shade(cell, BRAND if cols > 2 else BRAND_MID)
            else:
                add_inline(p, text, 9.5, TEXT)
                if ri % 2 == 0:
                    cell_shade(cell, SOFT2)
    repeat_header(t)
    fixed_layout(t, col_widths(rows))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_paragraph(doc: Document, text: str) -> None:
    label, rest = split_label(text)

    if label == "证据":
        p = doc.add_paragraph(style="DemoEvidence")
        para_border(p, "left", GREEN, 12, 6)
        para_shade(p, GREEN_SOFT)
        r = p.add_run("证据：")
        font_of(r, FONT, 9, True, "2F7A24")
        add_inline(p, rest, 9, DIM, code_color=BRAND_MID)
        return

    p = doc.add_paragraph(style="DemoBody")
    if label:
        r = p.add_run(label + "：")
        font_of(r, FONT, 10.5, True, ACCENT if label in ("操作", "操作顺序") else BRAND_MID)
        add_inline(p, rest, 10.5)
    else:
        add_inline(p, text, 10.5)


def build(blocks: list[tuple]) -> Document:
    doc = Document()
    build_styles(doc)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    build_header_footer(doc, sec)

    cp = doc.core_properties
    cp.title = "厦大统一门户 · 功能演示文档"
    cp.subject = "320 条校园入口统一收口——演示脚本、证据与边界"
    cp.author = "厦大统一门户项目"
    cp.comments = "由 tools/make_demo_docx.py 从 DEMO.md 生成"

    first_h1 = True
    for kind, payload in blocks:
        if kind == "h1":
            p = doc.add_paragraph(payload, style="DemoTitle")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if kind == "h2":
            p = doc.add_paragraph(payload, style="DemoH1")
            if first_h1:
                first_h1 = False
            else:
                p.paragraph_format.page_break_before = False
            para_border(p, "bottom", SOFT, 12, 4)
            continue
        if kind == "h3":
            p = doc.add_paragraph(payload, style="DemoH2")
            if "⭐" in payload or "★" in payload:
                for r in p.runs:
                    r.font.color.rgb = RGBColor.from_string(ACCENT)
            continue
        if kind == "p":
            add_paragraph(doc, payload)
            continue
        if kind == "quote":
            for idx, line in enumerate(payload):
                p = doc.add_paragraph(style="DemoQuote")
                para_shade(p, SOFT)
                para_border(p, "left", ACCENT, 16, 6)
                add_inline(p, line, 10.5, BRAND_MID)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if kind == "ul":
            for item in payload:
                p = doc.add_paragraph(style="DemoBullet")
                add_inline(p, "· " + item, 10.5)
            continue
        if kind == "ol":
            for item in payload:
                p = doc.add_paragraph(style="DemoNumber")
                add_inline(p, item, 10.5)
            continue
        if kind == "code":
            for line in payload:
                p = doc.add_paragraph(style="DemoCode")
                para_shade(p, SOFT2)
                r = p.add_run(line if line.strip() else " ")
                font_of(r, MONO, 9, False, "1F3A5F")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if kind == "table":
            add_table(doc, payload)
            continue
        if kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            para_border(p, "bottom", SOFT, 6, 2)
            continue

    # 结尾签名（居中、浅色）
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run("厦大统一门户 · 学生作品，非厦门大学官方产品 · 数据整理自各单位官网公开信息")
    font_of(r, FONT, 8.5, False, MUTE)
    return doc


def verify() -> int:
    """生成后自检：能不能被重新读出来、XML 是否合法、内容有没有漏。"""
    import zipfile
    import xml.etree.ElementTree as ET

    bad = 0
    with zipfile.ZipFile(DST) as z:
        names = z.namelist()
        broken = z.testzip()
        if broken:
            print("zip 损坏：%s" % broken)
            bad += 1
        for n in names:
            if n.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(n))
                except ET.ParseError as exc:
                    print("XML 不合法 %s：%s" % (n, exc))
                    bad += 1
        doc_xml = z.read("word/document.xml").decode("utf-8")
        for need in ("word/styles.xml", "word/header1.xml", "word/footer1.xml",
                     "word/media/", "docProps/core.xml"):
            if not any(n.startswith(need) for n in names):
                print("缺部件：%s" % need)
                bad += 1
        if "w:fldSimple" not in z.read("word/footer1.xml").decode("utf-8"):
            print("页脚缺页码域")
            bad += 1
        if "r:embed" not in z.read("word/header1.xml").decode("utf-8"):
            print("页眉缺校徽图片")
            bad += 1
        for bad_val in ('w:val="1"', 'w:val="0"'):
            if "vAlign" in doc_xml and bad_val in doc_xml.split("vAlign")[1][:24]:
                print("垂直居中写成了数字：%s" % bad_val)
                bad += 1

        # 子元素顺序：顺序错了 Word 会弹"需要修复"，这里逐个数一遍
        root = ET.fromstring(doc_xml)
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for parent_tag in ("pPr", "tcPr", "tblPr"):
            order = ORDER["w:" + parent_tag]
            for el in root.iter(w + parent_tag):
                idxs = []
                for child in el:
                    name = "w:" + child.tag.replace(w, "")
                    if name in order:
                        idxs.append(order.index(name))
                if idxs != sorted(idxs):
                    print("%s 子元素顺序不对：%s" % (parent_tag, idxs))
                    bad += 1
                    break

        # 内容类型声明 + 关系完整性：这两样缺了，Word 打开就报"需要修复"
        import posixpath
        ctypes_xml = z.read("[Content_Types].xml").decode("utf-8")
        defaults = {e.lower() for e in re.findall(r'Extension="([^"]+)"', ctypes_xml)}
        overrides = {p.lstrip("/") for p in re.findall(r'PartName="([^"]+)"', ctypes_xml)}
        for part in ("word/document.xml", "word/styles.xml", "word/header1.xml",
                     "word/footer1.xml", "docProps/core.xml", "word/media/image1.png"):
            if part in overrides:
                continue
            ext = part.rsplit(".", 1)[-1].lower()
            if ext not in defaults:
                print("[Content_Types].xml 没声明 %s" % part)
                bad += 1
        for part in ("word/document.xml", "word/header1.xml", "word/footer1.xml"):
            rels_path = "%s/_rels/%s.rels" % (posixpath.dirname(part), posixpath.basename(part))
            rels: dict[str, str] = {}
            if rels_path in names:
                for rel in ET.fromstring(z.read(rels_path)):
                    rels[rel.get("Id")] = rel.get("Target", "")
            for rid in sorted(set(re.findall(r'r:(?:id|embed)="([^"]+)"',
                                             z.read(part).decode("utf-8")))):
                if rid not in rels:
                    print("%s 引用了不存在的关系 %s" % (part, rid))
                    bad += 1
            for rid, target in rels.items():
                if target.startswith(("http:", "https:", "mailto:")):
                    continue
                resolved = posixpath.normpath(
                    posixpath.join(posixpath.dirname(part), target))
                if resolved not in names:
                    print("关系目标缺失：%s → %s" % (rels_path, resolved))
                    bad += 1

    doc = Document(str(DST))
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            text += "\n" + "\t".join(c.text for c in row.cells)
    for marker in ("**", "`"):
        if marker in text:
            print("正文残留 markdown 记号 %r：%d 处" % (marker, text.count(marker)))
            bad += 1
    must = ["厦大统一门户", "宿舍水管漏了", "27.5", "89 条", "320", "教务服务平台"]
    for m in must:
        if m not in text:
            print("缺少关键内容：%s" % m)
            bad += 1
    if len(doc.tables) != 6:
        print("表格数不对：%d（应为 6）" % len(doc.tables))
        bad += 1

    # 结构不变量：块 → 段落的数量必须一一对上，防止解析时悄悄丢内容
    blocks = parse(SRC.read_text(encoding="utf-8"))
    want = {
        "DemoH1": sum(1 for k, _ in blocks if k == "h2"),
        "DemoH2": sum(1 for k, _ in blocks if k == "h3"),
        "DemoBody": sum(1 for k, _ in blocks if k == "p"),
        "DemoNumber": sum(len(v) for k, v in blocks if k == "ol"),
        "DemoCode": sum(len(v) for k, v in blocks if k == "code"),
        "DemoTitle": sum(1 for k, _ in blocks if k == "h1"),
    }
    got: dict[str, int] = {}
    for p in doc.paragraphs:
        got[p.style.name] = got.get(p.style.name, 0) + 1
    for style, n in want.items():
        # 证据 段单独成样式，从 DemoBody 里刨掉
        if style == "DemoBody":
            n -= sum(1 for k, v in blocks if k == "p" and split_label(v)[0] == "证据")
        if got.get(style, 0) != n:
            print("段落数不符 %s：docx %d，应为 %d" % (style, got.get(style, 0), n))
            bad += 1

    print("自检：段落 %d｜表格 %d｜字符 %d｜部件 %d｜问题 %d"
          % (len(doc.paragraphs), len(doc.tables), len(text), len(names), bad))
    return bad


def main() -> int:
    md = SRC.read_text(encoding="utf-8")
    blocks = parse(md)
    kinds: dict[str, int] = {}
    for k, _ in blocks:
        kinds[k] = kinds.get(k, 0) + 1
    doc = build(blocks)
    doc.save(DST)
    size = DST.stat().st_size
    print("已生成 %s" % DST.name)
    print("大小 %.1f KB｜块：%s" % (size / 1024.0, kinds))
    print("段落 %d｜表格 %d" % (len(doc.paragraphs), len(doc.tables)))
    return 1 if verify() else 0


if __name__ == "__main__":
    raise SystemExit(main())
